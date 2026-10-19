from openpyxl import load_workbook

from loaders.base import BaseLoader
from models.document import Document


class XLSXLoader(BaseLoader):

    def load(
        self,
        file_path: str,
    ) -> Document:

        workbook = load_workbook(
            filename=file_path,
            data_only=True,
        )

        text_parts = []

        for sheet in workbook.worksheets:

            text_parts.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):

                values = []

                for cell in row:

                    if cell is None:
                        continue

                    values.append(str(cell))

                if values:
                    text_parts.append(" | ".join(values))

            text_parts.append("")

        return Document(
            text="\n".join(text_parts),
            page_count=len(workbook.sheetnames),
            metadata={
                "type": "xlsx",
                "sheet_count": len(workbook.sheetnames),
                "sheet_names": workbook.sheetnames,
            },
        )